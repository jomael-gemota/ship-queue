"""Read CSV / Excel invoice exports and split them into per-invoice row groups.

Exports usually hold many invoices per file — one or more rows each — so the
rows are grouped by invoice number (falling back to PO) and Tidy reads the
groups in batches (see design-log/2026-10-22-csv-excel-invoice-parsing.md).
"""

from __future__ import annotations

import csv
import io
import logging
import os
import re
from dataclasses import dataclass
from datetime import date, datetime, time

logger = logging.getLogger(__name__)

TABULAR_KINDS = ("csv", "xlsx", "xls")

BATCH_MAX_INVOICES = int(os.environ.get("TIDY_TABULAR_BATCH_INVOICES", 10))
BATCH_MAX_CHARS = int(os.environ.get("TIDY_TABULAR_BATCH_CHARS", 12_000))

# Only the opening rows are scanned for the header, so a title row or two above
# it is skipped but a long body never is.
_HEADER_SCAN_ROWS = 10

_INVOICE_KEY_HEADERS = {
    "invoice", "invoiceid", "invoicenumber", "invoiceno", "invoicenum", "invoicenbr",
    "inv", "invid", "invnumber", "invno", "invnum",
}
# Priority order: explicit PO headers before customer PO / reference ones.
_PO_KEY_HEADERS = (
    "ponumber", "po", "pono", "ponum", "purchaseordernumber", "purchaseorder",
    "customerponumber", "customerpo", "customerpono", "custponumber", "custpo",
    "customerpurchaseorder", "yourpo", "buyerpo",
    "customerreferencenumber", "customerreference", "customerref", "custref",
    "yourreference", "yourref",
)


@dataclass
class Spreadsheet:
    header: list[str]
    rows: list[list[str]]


def file_kind(filename: str | None) -> str:
    """`pdf`, `csv`, `xlsx` or `xls` from the filename; anything else reads as `pdf`."""
    ext = (filename or "").rsplit(".", 1)[-1].lower() if "." in (filename or "") else ""
    return ext if ext in TABULAR_KINDS else "pdf"


def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.lower())


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _read_csv(data: bytes) -> list[list[str]]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
            if sum(1 for r in rows if any(r)) >= 2:
                return rows
    finally:
        workbook.close()
    return []


def _read_xls(data: bytes) -> list[list[str]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    for sheet in book.sheets():
        rows: list[list[str]] = []
        for r in range(sheet.nrows):
            row: list[str] = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(_cell(xlrd.xldate_as_datetime(cell.value, book.datemode)))
                else:
                    row.append(_cell(cell.value))
            rows.append(row)
        if sum(1 for r in rows if any(r)) >= 2:
            return rows
    return []


def read_spreadsheet(data: bytes, kind: str) -> Spreadsheet:
    """Header + data rows of the first non-empty sheet, trailing blank columns trimmed."""
    reader = {"csv": _read_csv, "xlsx": _read_xlsx, "xls": _read_xls}[kind]
    rows = [r for r in reader(data) if any(c for c in r)]
    if not rows:
        raise ValueError("The spreadsheet has no data rows.")

    filled = [sum(1 for c in r if c) for r in rows[:_HEADER_SCAN_ROWS]]
    header_at = filled.index(max(filled))
    header = rows[header_at]
    body = rows[header_at + 1:]
    if not body:
        raise ValueError("The spreadsheet has a header row but no data rows.")

    width = max(len(r) for r in [header, *body])
    header = header + [""] * (width - len(header))
    body = [r + [""] * (width - len(r)) for r in body]
    while width > 1 and not header[width - 1] and all(not r[width - 1] for r in body):
        width -= 1
    return Spreadsheet(header=header[:width], rows=[r[:width] for r in body])


def _key_column(header: list[str]) -> int | None:
    normed = [_norm(h) for h in header]
    for keys in (_INVOICE_KEY_HEADERS, _PO_KEY_HEADERS):
        for i, h in enumerate(normed):
            if h in keys:
                return i
    return None


def po_value(sheet: Spreadsheet, rows: list[list[str]]) -> str | None:
    """The invoice's PO straight from its PO / customer reference column, if the sheet has one."""
    normed = [_norm(h) for h in sheet.header]
    for key in _PO_KEY_HEADERS:
        if key in normed:
            col = normed.index(key)
            return next((r[col] for r in rows if r[col]), None)
    return None


def group_invoices(sheet: Spreadsheet) -> list[list[list[str]]]:
    """Rows grouped per invoice, in first-appearance order.

    Returns a single group (the whole sheet) when no invoice or PO column exists,
    so such files are read as one document.
    """
    col = _key_column(sheet.header)
    if col is None:
        return [sheet.rows]

    groups: dict[str, list[list[str]]] = {}
    last_key: str | None = None
    for row in sheet.rows:
        key = row[col].strip() or last_key
        if key is None:
            key = ""
        groups.setdefault(key, []).append(row)
        last_key = key
    return list(groups.values())


def _render_rows(header: list[str], rows: list[list[str]]) -> str:
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(header)
    writer.writerows(rows)
    return out.getvalue().rstrip("\n")


def sheet_text(sheet: Spreadsheet) -> str:
    """The whole sheet as CSV text (vendor detection, correction retrieval, PO fallback)."""
    return _render_rows(sheet.header, sheet.rows)


def invoice_text(sheet: Spreadsheet, rows: list[list[str]]) -> str:
    return _render_rows(sheet.header, rows)


def batch_text(sheet: Spreadsheet, groups: list[list[list[str]]], first_number: int, total: int) -> str:
    """One block per invoice, each with the header repeated so columns can't drift."""
    blocks = [
        f"--- Invoice {first_number + i} of {total} ---\n{invoice_text(sheet, rows)}"
        for i, rows in enumerate(groups)
    ]
    return (
        f"Spreadsheet export — {len(groups)} invoice(s) below, one block each.\n\n"
        + "\n\n".join(blocks)
    )


def batch_groups(sheet: Spreadsheet, groups: list[list[list[str]]]) -> list[list[list[list[str]]]]:
    """Consecutive invoice groups packed into batches by count and size."""
    batches: list[list[list[list[str]]]] = []
    current: list[list[list[str]]] = []
    size = 0
    for rows in groups:
        length = len(invoice_text(sheet, rows))
        if current and (len(current) >= BATCH_MAX_INVOICES or size + length > BATCH_MAX_CHARS):
            batches.append(current)
            current, size = [], 0
        current.append(rows)
        size += length
    if current:
        batches.append(current)
    return batches


def _find_line_items(inv: dict) -> list[dict]:
    """Return the line-item array from an invoice dict, or [] if none is found.

    Checks common field names first; falls back to the first list-of-dicts field.
    """
    _LI_NORMS = {
        "lineitems", "lineitem", "items", "lines", "lineitemlist",
        "details", "products", "orderitems", "orderlines",
    }
    # Priority pass: recognised key names.
    for key, value in inv.items():
        if _norm(key) in _LI_NORMS and isinstance(value, list):
            return [i for i in value if isinstance(i, dict)]
    # Fallback: first list-of-dicts field regardless of name.
    for value in inv.values():
        if isinstance(value, list) and value and isinstance(value[0], dict):
            return [i for i in value if isinstance(i, dict)]
    return []


def _line_items_table(invoices: list[dict], summary_columns: list[str]) -> dict | None:
    """Combined read-only display table for all invoices' line items.

    Includes an invoice-identifier prefix column so each row can be traced back
    to its invoice.  Returns None when no invoice has line items.

    This table is intentionally skipped by ``tablesToAgentJson`` on the server
    when saving a tabular correction (the "Invoices" table is the only editable
    one for batch files), so it never creates a spurious top-level key.
    """
    # Pick the best invoice identifier from the scalar columns already collected.
    _ID_NORMS = {
        "invoicenumber", "invoiceno", "invoiceid", "invnumber", "invno", "invid",
        "ponumber", "po", "pono",
    }
    inv_id_col: str | None = next(
        (c for c in summary_columns if _norm(c) in _ID_NORMS),
        summary_columns[0] if summary_columns else None,
    )

    li_columns: list[str] = []  # column names from line-item dicts
    collected: list[tuple[str, dict]] = []  # (inv_id_value, item_dict)

    for inv in invoices:
        items = _find_line_items(inv)
        if not items:
            continue
        inv_id_val = str(inv.get(inv_id_col, "") if inv_id_col else "")
        for item in items:
            for k, v in item.items():
                if k not in li_columns and not isinstance(v, (dict, list)):
                    li_columns.append(k)
            collected.append((inv_id_val, item))

    if not collected:
        return None

    # Prepend the invoice identifier only when it isn't already a line-item key.
    if inv_id_col and inv_id_col not in li_columns:
        all_columns = [inv_id_col] + li_columns
        def make_row(inv_id_val: str, item: dict) -> list:
            return [inv_id_val] + [
                "" if item.get(c) is None else item.get(c, "")
                for c in li_columns
            ]
    else:
        all_columns = li_columns
        def make_row(inv_id_val: str, item: dict) -> list:  # noqa: F811
            return ["" if item.get(c) is None else item.get(c, "") for c in li_columns]

    rows = [make_row(inv_id_val, item) for inv_id_val, item in collected]
    return {"title": "Line Items", "columns": all_columns, "rows": rows}


def invoice_table(invoices: list[dict]) -> dict:
    """Display table for a batch file.

    Always emits an ``"Invoices"`` summary grid (one row per invoice, scalar
    fields only) that round-trips cleanly through tabular corrections.

    When invoices carry line items, also emits a combined **read-only**
    ``"Line Items"`` grid so all CSV columns are visible in the Table view.
    ``tablesToAgentJson`` on the server skips this table during corrections
    (it has no matching top-level key in the batch JSON structure).
    """
    # ── Invoices summary table (editable via tabular correction) ──
    columns: list[str] = []
    for inv in invoices:
        for key, value in inv.items():
            if key not in columns and not isinstance(value, (dict, list)):
                columns.append(key)
    rows = [
        [inv.get(c) if not isinstance(inv.get(c), (dict, list)) else "" for c in columns]
        for inv in invoices
    ]
    rows = [["" if v is None else v for v in row] for row in rows]
    tables: list[dict] = [{"title": "Invoices", "columns": columns, "rows": rows}]

    # ── Line items table (read-only display supplement) ──
    li_table = _line_items_table(invoices, columns)
    if li_table is not None:
        tables.append(li_table)

    return {"tables": tables}
